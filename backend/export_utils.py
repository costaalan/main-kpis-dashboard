"""Export modules: PPT and HTML report generation from chat history + KPI data."""

import io
from datetime import datetime
from pathlib import Path

import cache_db
import agent as agent_mod


def format_value(valor, kpi_name):
    """Format a KPI value for display."""
    if kpi_name in ("GMV",):
        if abs(valor) >= 1_000_000:
            return f"R$ {valor/1_000_000:,.2f}M"
        elif abs(valor) >= 1_000:
            return f"R$ {valor/1_000:,.1f}K"
        return f"R$ {valor:,.2f}"
    elif kpi_name in ("AOV", "CAC"):
        return f"R$ {valor:,.2f}"
    elif kpi_name in ("CVR", "ROAS", "OTD", "TAXA_RECLAMACAO", "TAXA_CANCELAMENTO"):
        return f"{valor*100:.2f}%"
    elif kpi_name == "TEMPO_DESPACHO":
        return f"{valor:.1f} dias"
    elif kpi_name == "LTV":
        return f"R$ {valor:,.2f}"
    return f"{valor:,.2f}"


def build_chart_data(conn, site_id, kpi_name, periodicidade):
    """Fetch data and build chart structure for a single KPI."""
    rows = cache_db.list_series(conn, site_id, kpi_name, periodicidade)
    if not rows:
        return None
    return {
        "kpi_name": kpi_name,
        "site_id": site_id,
        "site_label": agent_mod.SITE_LABELS.get(site_id, site_id),
        "periodicidade": periodicidade,
        "rows": rows,
        "formatted": [{"label": r["periodo_label"], "valor": format_value(r["valor"], kpi_name)} for r in rows],
    }


def generate_html_report(conn, chat_messages, site_id=None):
    """Generate a standalone HTML report from chat history and KPI data."""
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    
    # Extract KPI data mentions from chat
    kpi_data_sections = []
    for msg in chat_messages:
        if msg.get("chart"):
            c = msg["chart"]
            kpi_data_sections.append(build_chart_data(
                conn, c["site_id"], c["kpi_name"], c.get("periodicidade", "dia")
            ))
    
    # Build KPI data HTML
    kpi_html = ""
    for data in kpi_data_sections:
        if not data:
            continue
        rows_html = "\n".join(
            f"<tr><td>{r['label']}</td><td style='text-align:right'>{r['valor']}</td></tr>"
            for r in data["formatted"][-10:]  # last 10 periods
        )
        kpi_html += f"""
        <div class="kpi-block">
            <h3>{data['kpi_name']} — {data['site_label']} ({data['site_id']})</h3>
            <table>
                <thead><tr><th>Periodo</th><th>Valor</th></tr></thead>
                <tbody>{rows_html}</tbody>
            </table>
        </div>"""
    
    # Build chat transcript
    chat_html = ""
    for msg in chat_messages:
        role_class = "user" if msg["role"] == "user" else "assistant"
        content = msg.get("content", "").replace("\n", "<br>")
        chat_html += f'<div class="msg {role_class}"><strong>{role_class.title()}:</strong><br>{content}</div>\n'
    
    return f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Relatorio KPIs — Main KPIs Dashboard</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; max-width: 900px; margin: 0 auto; padding: 20px; background: #f8fafc; color: #1e293b; }}
  h1 {{ color: #2563eb; border-bottom: 3px solid #2563eb; padding-bottom: 10px; }}
  h2 {{ color: #334155; margin-top: 30px; }}
  h3 {{ color: #475569; }}
  .meta {{ color: #94a3b8; font-size: 0.85em; margin-bottom: 20px; }}
  .section {{ background: white; border-radius: 8px; padding: 20px; margin: 20px 0; box-shadow: 0 1px 3px rgba(0,0,0,0.1); }}
  .kpi-block {{ margin: 15px 0; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 0.9em; }}
  th, td {{ padding: 8px 12px; border-bottom: 1px solid #e2e8f0; text-align: left; }}
  th {{ background: #f1f5f9; font-weight: 600; }}
  tr:nth-child(even) {{ background: #f8fafc; }}
  .chat {{ margin-top: 30px; }}
  .msg {{ padding: 10px 15px; margin: 8px 0; border-radius: 8px; }}
  .msg.user {{ background: #dbeafe; text-align: right; }}
  .msg.assistant {{ background: #f1f5f9; text-align: left; }}
  .footer {{ text-align: center; color: #94a3b8; font-size: 0.8em; margin-top: 40px; }}
  @media print {{ body {{ background: white; }} .section {{ box-shadow: none; border: 1px solid #e2e8f0; }} }}
</style>
</head>
<body>
<h1>Relatorio de KPIs</h1>
<p class="meta">Gerado em {now} — Main KPIs Dashboard</p>

<div class="section">
  <h2>Dados dos KPIs</h2>
  {kpi_html if kpi_html else '<p>Nenhum dado de KPI mencionado na conversa.</p>'}
</div>

<div class="section chat">
  <h2>Transcricao do Chat</h2>
  {chat_html}
</div>

<div class="footer">
  <p>Main KPIs Dashboard — Relatorio gerado automaticamente</p>
</div>
</body>
</html>"""


def generate_ppt(conn, chat_messages, site_id=None):
    """Generate a PowerPoint presentation from chat history and KPI data."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN
        from pptx.dml.color import RGBColor
    except ImportError:
        raise RuntimeError("python-pptx not installed. Run: pip install python-pptx")
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    left, top, width, height = Inches(1), Inches(2), Inches(11), Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Relatorio de KPIs"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    top2 = Inches(3.8)
    txBox2 = slide.shapes.add_textbox(left, top2, width, Inches(1))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = f"Main KPIs Dashboard — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    p2.alignment = PP_ALIGN.CENTER
    
    # Extract KPI data
    kpi_data_sections = []
    for msg in chat_messages:
        if msg.get("chart"):
            c = msg["chart"]
            kpi_data_sections.append(build_chart_data(
                conn, c["site_id"], c["kpi_name"], c.get("periodicidade", "dia")
            ))
    
    # KPI data slide
    if kpi_data_sections:
        for data in kpi_data_sections:
            if not data:
                continue
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            
            # Title
            left, top = Inches(0.8), Inches(0.5)
            txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(0.8))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = f"{data['kpi_name']} — {data['site_label']} ({data['site_id']})"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            
            # Table
            rows_display = data["formatted"][-12:]
            table_top = Inches(1.8)
            table_width = Inches(11)
            table_height = Inches(0.4 * (len(rows_display) + 1))
            
            table_shape = slide.shapes.add_table(
                len(rows_display) + 1, 2, left, table_top, table_width, table_height
            )
            table = table_shape.table
            
            # Headers
            for col_idx, header in enumerate(["Periodo", "Valor"]):
                cell = table.cell(0, col_idx)
                cell.text = header
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(14)
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x25, 0x63, 0xEB)
            
            # Data rows
            for row_idx, row in enumerate(rows_display):
                for col_idx, key in enumerate(["label", "valor"]):
                    cell = table.cell(row_idx + 1, col_idx)
                    cell.text = row[key]
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.size = Pt(12)
                if row_idx % 2 == 0:
                    for col_idx in range(2):
                        table.cell(row_idx + 1, col_idx).fill.solid()
                        table.cell(row_idx + 1, col_idx).fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
    
    # Chat transcript slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    left, top = Inches(0.8), Inches(0.5)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Transcricao do Chat"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    # Chat content
    chat_top = Inches(1.8)
    chat_box = slide.shapes.add_textbox(left, chat_top, Inches(11), Inches(5))
    tf = chat_box.text_frame
    tf.word_wrap = True
    
    for i, msg in enumerate(chat_messages):
        if i > 0:
            p = tf.add_paragraph()
            p.space_before = Pt(8)
        else:
            p = tf.paragraphs[0]
        
        role = "Usuario" if msg["role"] == "user" else "Assistente"
        content = msg.get("content", "")[:300]  # Truncate long messages
        p.text = f"{role}: {content}"
        p.font.size = Pt(11)
        if msg["role"] == "user":
            p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        else:
            p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Save to bytes
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
